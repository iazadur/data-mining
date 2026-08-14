import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os
import re

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_docx_report():
    doc = docx.Document()
    
    # Page Margins (1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Style definitions
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    with open('LAB_REPORT.md', 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(lines_list):
        if not lines_list:
            return
        rows_data = []
        for l in lines_list:
            l_str = l.strip()
            if l_str.startswith('|') and l_str.endswith('|'):
                parts = [p.strip() for p in l_str.split('|')[1:-1]]
                # Ignore delimiter row (e.g. | :--- | :--- |)
                if any('---' in p for p in parts):
                    continue
                rows_data.append(parts)

        if not rows_data:
            return

        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for i, row in enumerate(rows_data):
            for j, val in enumerate(row):
                if j < len(table.columns):
                    cell = table.cell(i, j)
                    # Strip bold markers
                    clean_val = val.replace('**', '').replace('`', '')
                    cell.text = clean_val
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)

                    if i == 0:
                        set_cell_background(cell, '2B6CB0')
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.font.size = Pt(10)
                    else:
                        if i % 2 == 0:
                            set_cell_background(cell, 'F7FAFC')
                        for run in p.runs:
                            run.font.size = Pt(9.5)
        
        doc.add_paragraph() # Spacing after table

    def flush_code(lines_list):
        if not lines_list:
            return
        code_text = "".join(lines_list)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.2)

        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)

    for line in lines:
        raw_line = line.rstrip('\n')
        stripped = raw_line.strip()

        # Handle Code Block
        if stripped.startswith('```'):
            if in_code_block:
                flush_code(code_lines)
                code_lines = []
                in_code_block = False
            else:
                if in_table:
                    flush_table(table_lines)
                    table_lines = []
                    in_table = False
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(raw_line + '\n')
            continue

        # Handle Table Lines
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        else:
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False

        if not stripped:
            continue

        # Handle Images ![alt](images/filename.png)
        img_match = re.match(r'!\[.*?\]\((images/.*?)\)', stripped)
        if img_match:
            img_path = img_match.group(1)
            if os.path.exists(img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(img_path, width=Inches(5.5))
                p.paragraph_format.space_after = Pt(12)
            continue

        # Handle Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(stripped[2:].replace('**', ''))
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[3:].replace('**', ''))
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[4:].replace('**', ''))
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        elif stripped.startswith('---'):
            continue
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            # Simple bold parsing
            text = stripped[2:]
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    r = p.add_run(part[1:-1])
                    r.font.italic = True
                elif part.startswith('`') and part.endswith('`'):
                    r = p.add_run(part[1:-1])
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9.5)
                else:
                    p.add_run(part)

    # Flush remaining table if exists
    if in_table:
        flush_table(table_lines)

    doc.save('LAB_REPORT.docx')
    print("DOCX successfully generated: LAB_REPORT.docx")

if __name__ == '__main__':
    create_docx_report()
